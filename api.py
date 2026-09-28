"""REST API for the React frontend. See docs/api-contract.md.

Run locally:  uvicorn api:app --reload --port 8000

Settings (environment variables):
  ANTHROPIC_API_KEY, DATABASE_URL, OWNER_EMAILS   as for the Streamlit app
  SECRET_KEY            signs sign-in tokens and session cookies (long random string)
  GOOGLE_CLIENT_ID, GOOGLE_CLIENT_SECRET          Google sign-in
  API_BASE_URL          this API's public URL, for Google's redirect (e.g. https://api.example.com)
  FRONTEND_ORIGINS      comma-separated frontend URLs allowed to call the API and receive sign-in redirects
  FRONTEND_ORIGIN_REGEX optional pattern for more frontend URLs (e.g. this project's Lovable preview addresses)
  DEV_LOGIN=1           local development only: enables /api/auth/dev-login without Google
"""

import io
import os
import re
import threading
import uuid
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Callable
from urllib.parse import urlparse

import docx
from authlib.integrations.starlette_client import OAuth
from dotenv import load_dotenv
from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, RedirectResponse
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from pydantic import BaseModel
from starlette.middleware.sessions import SessionMiddleware

load_dotenv(Path(__file__).parent / ".env")

import anthropic  # noqa: E402  (after load_dotenv so the client sees the key)

import coaching  # noqa: E402
import database  # noqa: E402
import job_sources  # noqa: E402
import limits  # noqa: E402
import logos  # noqa: E402
import plans  # noqa: E402
import referrals  # noqa: E402
import scout  # noqa: E402
from analyzer import MODEL, AnalysisError, JobFitAnalysis, analyze_fit  # noqa: E402
from discovery import discover_for_user  # noqa: E402
from search_profile import SearchProfile  # noqa: E402

SECRET_KEY = os.getenv("SECRET_KEY", "")
if not SECRET_KEY:
    if os.getenv("DEV_LOGIN") != "1":
        raise RuntimeError("Set SECRET_KEY (a long random string) before starting the API.")
    SECRET_KEY = "local-development-only"
FRONTEND_ORIGINS = [o.strip().rstrip("/") for o in os.getenv("FRONTEND_ORIGINS", "").split(",") if o.strip()]
TOKEN_MAX_AGE = 30 * 24 * 3600
# Only our own frontends may call the API or receive sign-in tokens. Never allow a whole
# hosting domain like *.lovable.app: anyone can publish a site there.
FRONTEND_ORIGIN_REGEX = os.getenv("FRONTEND_ORIGIN_REGEX") or None

app = FastAPI(title="Job Search Copilot API")
app.add_middleware(SessionMiddleware, secret_key=SECRET_KEY, same_site="lax", https_only=False)
app.add_middleware(
    CORSMiddleware,
    allow_origins=FRONTEND_ORIGINS,
    allow_origin_regex=FRONTEND_ORIGIN_REGEX,
    allow_methods=["*"],
    allow_headers=["*"],
)
database.init_db()
signer = URLSafeTimedSerializer(SECRET_KEY, salt="auth")


@app.exception_handler(limits.LimitReached)
def _limit(_: Request, e: limits.LimitReached):
    return JSONResponse({"detail": str(e)}, status_code=429)


@app.exception_handler(anthropic.APIError)
def _ai_error(_: Request, e: anthropic.APIError):
    return JSONResponse({"detail": "The AI service had a problem. Please try again in a moment."}, status_code=502)


@app.exception_handler(AnalysisError)
def _analysis_error(_: Request, e: AnalysisError):
    return JSONResponse({"detail": str(e)}, status_code=502)


# --- Sign-in ---------------------------------------------------------------------------------

oauth = OAuth()
oauth.register(
    "google",
    client_id=os.getenv("GOOGLE_CLIENT_ID"),
    client_secret=os.getenv("GOOGLE_CLIENT_SECRET"),
    server_metadata_url="https://accounts.google.com/.well-known/openid-configuration",
    client_kwargs={"scope": "openid email profile", "prompt": "select_account"},
)


def _allowed_redirect(url: str) -> bool:
    """Only send sign-in tokens back to our own frontends."""
    origin = "{0.scheme}://{0.netloc}".format(urlparse(url))
    return origin in FRONTEND_ORIGINS or bool(FRONTEND_ORIGIN_REGEX and re.fullmatch(FRONTEND_ORIGIN_REGEX, origin))


@app.get("/api/auth/login")
async def login(request: Request, redirect: str):
    if not _allowed_redirect(redirect):
        raise HTTPException(400, "That redirect address isn't allowed.")
    request.session["redirect"] = redirect
    callback = os.getenv("API_BASE_URL", str(request.base_url).rstrip("/")) + "/api/auth/callback"
    return await oauth.google.authorize_redirect(request, callback)


@app.get("/api/auth/callback")
async def auth_callback(request: Request):
    token = await oauth.google.authorize_access_token(request)
    info = token.get("userinfo") or {}
    if not info.get("email") or not info.get("email_verified", True):
        raise HTTPException(400, "Google didn't confirm an email address for this account.")
    user = database.get_or_create_user(info["email"], info.get("name", ""))
    redirect = request.session.pop("redirect", FRONTEND_ORIGINS[0] if FRONTEND_ORIGINS else "/")
    return RedirectResponse(f"{redirect}#token={signer.dumps(user['id'])}")


@app.get("/api/auth/dev-login")
def dev_login(redirect: str = ""):
    """Local development only (DEV_LOGIN=1): sign in as the local user without Google."""
    if os.getenv("DEV_LOGIN") != "1":
        raise HTTPException(404)
    user = database.get_or_create_user(limits.LOCAL_USER_EMAIL, "Local user")
    token = signer.dumps(user["id"])
    return RedirectResponse(f"{redirect}#token={token}") if redirect else {"token": token}


def current_user(authorization: str = Header("")) -> dict:
    token = authorization.removeprefix("Bearer ").strip()
    try:
        user_id = signer.loads(token, max_age=TOKEN_MAX_AGE)
    except SignatureExpired:
        raise HTTPException(401, "Your session expired. Please sign in again.")
    except BadSignature:
        raise HTTPException(401, "Please sign in.")
    user = database.get_user(user_id)
    if not user:
        raise HTTPException(401, "Please sign in.")
    return user


# --- Background jobs -------------------------------------------------------------------------
# Long actions (finding roles, discovery, skills reports, analyses) run in a background
# thread; the frontend polls GET /api/jobs/{id}, which also carries progress messages.

JOBS: dict[str, dict] = {}


def start_job(user: dict, work: Callable[[Callable[[str], None]], Any]) -> dict:
    job_id = uuid.uuid4().hex
    job = {"id": job_id, "user_id": user["id"], "status": "running", "log": [], "result": None, "error": None}
    JOBS[job_id] = job

    def run():
        try:
            job["result"] = work(job["log"].append)
            job["status"] = "done"
        except limits.LimitReached as e:
            job["error"], job["status"] = str(e), "error"
        except anthropic.APIError:
            job["error"], job["status"] = "The AI service had a problem. Please try again in a moment.", "error"
        except Exception as e:  # show a readable message; details go to the server log
            print(f"Job {job_id} failed: {e!r}")
            job["error"], job["status"] = str(e) or "Something went wrong.", "error"

    threading.Thread(target=run, daemon=True).start()
    return {"job_id": job_id}


@app.get("/api/jobs/{job_id}")
def get_job(job_id: str, user: dict = Depends(current_user)):
    job = JOBS.get(job_id)
    if not job or job["user_id"] != user["id"]:
        raise HTTPException(404, "That task wasn't found. It may have been interrupted; please try again.")
    return {k: job[k] for k in ("status", "log", "result", "error")}


# --- Shapes ----------------------------------------------------------------------------------

def iso(value) -> str:
    """Timestamps as ISO 8601, which every browser can parse."""
    return value.isoformat() if hasattr(value, "isoformat") else str(value)


def me_out(user: dict) -> dict:
    user = database.get_user(user["id"])
    usage = {kind: {"used": limits.daily_limit(user, kind) - limits.remaining(user, kind),
                    "limit": limits.daily_limit(user, kind)} for kind in limits.LIMITS}
    return {
        "name": user["name"], "email": user["email"], "is_owner": limits.is_owner(user),
        "onboarding": {"has_resume": bool(user["resume_pdf"] or user["resume_text"]),
                       "has_search": bool(user["profile_json"]),
                       "has_companies": bool(database.list_user_companies(user["id"]))},
        "resume_kind": "pdf" if user["resume_pdf"] else "text" if user["resume_text"] else None,
        "connections_count": len(database.list_connections(user["id"])),
        "last_scan": user["last_scan"], "usage": usage,
    }


def role_out(p: dict, connections: list[dict]) -> dict:
    label, color = plans.fit_label(p["fit_score"] or 0)
    return {
        "id": p["id"], "title": p["title"], "company": p["company"], "company_id": p["company_id"],
        "logo_url": p["logo_url"] or None, "location": p["location"], "posted_at": p["posted_at"], "url": p["url"],
        "fit_score": p["fit_score"] or 0, "fit_color": color, "fit_label": label, "fit_reason": p["fit_reason"] or "",
        "status": p["status"], "has_plan": bool(p["analysis_id"]),
        "known_people": sum(referrals.same_company(c["company"], p["company"]) for c in connections),
    }


def role_detail(user: dict, role_id: int) -> dict:
    p = database.get_user_posting(user["id"], role_id)
    if not p:
        raise HTTPException(404, "That role isn't in your list anymore.")
    connections = database.list_connections(user["id"])
    at_company = [c for c in connections if referrals.same_company(c["company"], p["company"])]
    contacts = referrals.rank_contacts(at_company, p["title"])
    analysis = database.get_analysis(user["id"], p["analysis_id"]) if p["analysis_id"] else None
    out = {"role": role_out(p, connections), "analysis": None, "recommendation": None, "plan": [],
           "people": [{"index": i, "name": c.name, "position": c.position, "url": c.url, "why": c.why}
                      for i, c in enumerate(contacts)],
           "has_connections": bool(connections), "adjacent": []}
    if not analysis:
        return out
    done = database.get_plan_done(p)
    rec = plans.recommendation(analysis.match_score, p["company"])
    profile = database.get_profile(user)
    watched = [c["company_id"] for c in database.list_user_companies(user["id"], "tracking") if c["company_id"]]
    openings = database.open_postings_at(list({*watched, p["company_id"]}))
    out.update({
        "analysis": analysis.model_dump(),
        "recommendation": {"headline": rec.headline, "detail": rec.detail, "color": rec.color},
        "plan": [{"key": s.key, "kind": s.kind, "title": s.title, "detail": s.detail, "done": s.key in done}
                 for s in plans.build_plan(analysis, p["company"], contacts, bool(connections))],
        "adjacent": [{"title": title, "why": why,
                      "in_search": title.lower() in {t.lower() for t in profile.include_titles},
                      "openings": [{"title": m["title"], "company": m["company"], "logo_url": m["logo_url"] or None,
                                    "location": m["location"], "url": m["url"]} for m in matches]}
                     for title, why, matches in plans.find_adjacent_openings(analysis, p, openings)],
    })
    return out


def companies_out(user: dict) -> dict:
    ensure_logos(user)
    rows = database.list_user_companies(user["id"])
    role_counts = Counter(p["company_id"] for p in database.list_scored_postings(user["id"], 0))
    connections = database.list_connections(user["id"])

    def one(c):
        return {"id": c["id"], "company_id": c["company_id"], "name": c["name"], "logo_url": c["logo_url"] or None,
                "board_name": {"greenhouse": "Greenhouse", "lever": "Lever", "ashby": "Ashby"}.get(c["ats"]),
                "board_url": job_sources.board_page_url(c["ats"], c["ats_slug"]) if c["ats"] else None,
                "why_it_fits": c["why_it_fits"], "open_roles": role_counts.get(c["company_id"], 0),
                "known_people": sum(referrals.same_company(k["company"], c["name"]) for k in connections)}

    return {status_key: [one(c) for c in rows if c["status"] == status]
            for status_key, status in [("suggested", "suggested"), ("watching", "tracking"),
                                       ("untrackable", "unverified")]}


def ensure_logos(user: dict) -> None:
    ids = [c["company_id"] for c in database.list_user_companies(user["id"]) if c["company_id"]]
    for company in database.companies_missing_logo(ids):
        domain, url = logos.find_logo(company["name"], company["website"] or "")
        database.set_company_logo(company["id"], domain, url)


# --- Account, resume, search -------------------------------------------------------------------

@app.get("/api/me")
def get_me(user: dict = Depends(current_user)):
    return me_out(user)


@app.delete("/api/me")
def delete_me(user: dict = Depends(current_user)):
    database.delete_user(user["id"])
    return {}


@app.put("/api/resume")
async def put_resume(file: UploadFile | None = File(None), text: str | None = Form(None),
                     user: dict = Depends(current_user)):
    if file is not None:
        data = await file.read()
        name = (file.filename or "").lower()
        if name.endswith(".pdf"):
            database.save_resume(user["id"], pdf=data)
        elif name.endswith(".docx"):
            database.save_resume(user["id"], text=_docx_text(data))
        elif name.endswith((".txt", ".md")):
            database.save_resume(user["id"], text=data.decode("utf-8", errors="ignore"))
        else:
            raise HTTPException(400, "Upload your resume as a PDF, Word (.docx) or text file.")
    elif text and text.strip():
        database.save_resume(user["id"], text=text)
    else:
        raise HTTPException(400, "Upload a resume file or paste its text.")
    return me_out(user)


def _docx_text(data: bytes) -> str:
    """Plain text from a Word document, including text inside tables."""
    try:
        doc = docx.Document(io.BytesIO(data))
    except Exception:
        raise HTTPException(400, "That Word file couldn't be read. Try saving it as a PDF.")
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    text = "\n".join(line for line in lines if line.strip())
    if not text.strip():
        raise HTTPException(400, "That Word file looks empty. Try saving it as a PDF.")
    return text


@app.get("/api/search")
def get_search(user: dict = Depends(current_user)):
    return database.get_profile(user).model_dump()


@app.put("/api/search")
def put_search(profile: SearchProfile, user: dict = Depends(current_user)):
    database.save_profile(user["id"], profile)
    return profile.model_dump()


# --- Roles ---------------------------------------------------------------------------------

def _posted(p: dict) -> datetime | None:
    try:
        posted = datetime.fromisoformat(p["posted_at"]) if p["posted_at"] else None
    except ValueError:
        return None
    return posted.replace(tzinfo=timezone.utc) if posted and posted.tzinfo is None else posted


@app.get("/api/roles")
def get_roles(view: str = "best", posted_within: int | None = None, sort: str = "fit",
              company_id: int | None = None, user: dict = Depends(current_user)):
    """posted_within: only roles posted in the last N days (roles without a date are hidden).
    sort: "fit" (best fit first) or "recent" (newest first).
    company_id: only roles at this company (Company.company_id)."""
    ensure_logos(user)
    profile = database.get_profile(user)
    connections = database.list_connections(user["id"])
    all_roles = database.list_scored_postings(user["id"], 0, ("new", "saved"))
    if view == "saved":
        roles = [p for p in all_roles if p["status"] == "saved"]
    elif view == "all":
        roles = all_roles
    else:
        roles = [p for p in all_roles if p["fit_score"] >= profile.min_score]
    if company_id:
        roles = [p for p in roles if p["company_id"] == company_id]
    if posted_within:
        cutoff = datetime.now(timezone.utc) - timedelta(days=posted_within)
        roles = [p for p in roles if (posted := _posted(p)) and posted >= cutoff]
    if sort == "recent":
        oldest = datetime.min.replace(tzinfo=timezone.utc)
        roles.sort(key=lambda p: (_posted(p) or oldest, p["fit_score"]), reverse=True)
    return {"stats": {"strong_count": sum(p["fit_score"] >= profile.min_score for p in all_roles),
                      "saved_count": sum(p["status"] == "saved" for p in all_roles),
                      "companies_watched": len(database.list_user_companies(user["id"], "tracking")),
                      "min_score": profile.min_score},
            "roles": [role_out(p, connections) for p in roles]}


@app.post("/api/roles/refresh")
def refresh_roles(user: dict = Depends(current_user)):
    def work(log):
        return {"summary": scout.run_scan([user["id"]], log=log, auto_analyze=False)[user["id"]]}
    return start_job(user, work)


@app.post("/api/roles/find")
def find_roles(user: dict = Depends(current_user)):
    def work(log):
        if not database.list_user_companies(user["id"], "suggested"):
            log("Researching companies that match your search and background…")
            discover_for_user(user, log=log)
        for c in database.list_user_companies(user["id"], "suggested"):
            database.set_user_company_status(user["id"], c["id"], "tracking")
        log("Scanning their job boards and scoring each role against your resume…")
        return {"summary": scout.run_scan([user["id"]], log=log, auto_analyze=False)[user["id"]]}
    return start_job(user, work)


@app.get("/api/roles/{role_id}")
def get_role(role_id: int, user: dict = Depends(current_user)):
    return role_detail(user, role_id)


class StatusIn(BaseModel):
    status: str


@app.post("/api/roles/{role_id}/status")
def set_role_status(role_id: int, body: StatusIn, user: dict = Depends(current_user)):
    if body.status not in ("new", "saved", "dismissed"):
        raise HTTPException(400, "Unknown status.")
    database.set_posting_status(user["id"], role_id, body.status)
    return role_detail(user, role_id)["role"]


@app.post("/api/roles/{role_id}/plan")
def build_plan(role_id: int, user: dict = Depends(current_user)):
    p = database.get_user_posting(user["id"], role_id)
    if not p:
        raise HTTPException(404, "That role isn't in your list anymore.")
    limits.require(user, "analyses")

    def work(log):
        log("Reading the job description and comparing it to your resume…")
        scout.analyze_posting(database.get_user(user["id"]), p)
        return role_detail(user, role_id)
    return start_job(user, work)


class StepIn(BaseModel):
    done: bool


@app.put("/api/roles/{role_id}/plan/{step_key}")
def set_plan_step(role_id: int, step_key: str, body: StepIn, user: dict = Depends(current_user)):
    p = database.get_user_posting(user["id"], role_id)
    if not p:
        raise HTTPException(404, "That role isn't in your list anymore.")
    done = database.get_plan_done(p)
    done = done | {step_key} if body.done else done - {step_key}
    database.set_plan_done(user["id"], role_id, done)
    return role_detail(user, role_id)["plan"]


class DraftIn(BaseModel):
    person_index: int


@app.post("/api/roles/{role_id}/draft")
def draft(role_id: int, body: DraftIn, user: dict = Depends(current_user)):
    p = database.get_user_posting(user["id"], role_id)
    if not p:
        raise HTTPException(404, "That role isn't in your list anymore.")
    at_company = [c for c in database.list_connections(user["id"]) if referrals.same_company(c["company"], p["company"])]
    contacts = referrals.rank_contacts(at_company, p["title"])
    if not 0 <= body.person_index < len(contacts):
        raise HTTPException(404, "That person wasn't found.")
    limits.require(user, "messages")
    analysis = database.get_analysis(user["id"], p["analysis_id"]) if p["analysis_id"] else None
    message = referrals.draft_message(contacts[body.person_index], p["title"], p["company"], analysis,
                                      user["resume_text"], user["resume_pdf"])
    limits.use(user, "messages")
    return {"message": message}


class TitleIn(BaseModel):
    title: str


@app.post("/api/roles/{role_id}/add-title")
def add_title(role_id: int, body: TitleIn, user: dict = Depends(current_user)):
    profile = database.get_profile(user)
    if body.title.lower() not in {t.lower() for t in profile.include_titles}:
        profile = profile.model_copy(update={"include_titles": [*profile.include_titles, body.title]})
        database.save_profile(user["id"], profile)
    return profile.model_dump()


# --- Companies -------------------------------------------------------------------------------

@app.get("/api/companies")
def get_companies(user: dict = Depends(current_user)):
    return companies_out(user)


@app.post("/api/companies/discover")
def discover(user: dict = Depends(current_user)):
    limits.require(user, "discoveries")

    def work(log):
        discover_for_user(database.get_user(user["id"]), log=log)
        return companies_out(user)
    return start_job(user, work)


class CompanyIn(BaseModel):
    name: str
    careers_url: str = ""


@app.post("/api/companies")
def add_company(body: CompanyIn, user: dict = Depends(current_user)):
    name = body.name.strip()
    board = job_sources.find_board(name, body.careers_url.strip())
    if not board:
        raise HTTPException(404, f"Couldn't find a job board for {name}. We support companies hiring through "
                                 "Greenhouse, Lever or Ashby. Try pasting a link to one of their job postings.")
    company_id = database.upsert_company(name, *board, website=body.careers_url.strip())
    if not database.add_user_company(user["id"], name, "", company_id, "tracking"):
        raise HTTPException(409, f"{name} is already on your list.")
    return next(c for c in companies_out(user)["watching"] if c["name"].lower() == name.lower())


@app.post("/api/companies/{row_id}/status")
def set_company_status(row_id: int, body: StatusIn, user: dict = Depends(current_user)):
    if body.status not in ("tracking", "rejected"):
        raise HTTPException(400, "Unknown status.")
    database.set_user_company_status(user["id"], row_id, body.status)
    return companies_out(user)


@app.delete("/api/companies/{row_id}")
def remove_company(row_id: int, user: dict = Depends(current_user)):
    database.delete_user_company(user["id"], row_id)
    return companies_out(user)


# --- Resume match and saved analyses -----------------------------------------------------------

class MatchIn(BaseModel):
    job_description: str


@app.post("/api/match")
def match(body: MatchIn, user: dict = Depends(current_user)):
    if not body.job_description.strip():
        raise HTTPException(400, "Paste a job description first.")
    limits.require(user, "analyses")

    def work(log):
        log("Comparing your experience to the role…")
        u = database.get_user(user["id"])
        result = analyze_fit(body.job_description, resume_text=u["resume_text"], resume_pdf=u["resume_pdf"])
        limits.use(u, "analyses")
        database.save_analysis(u["id"], body.job_description, result, MODEL)
        return result.model_dump()
    return start_job(user, work)


@app.get("/api/analyses")
def get_analyses(user: dict = Depends(current_user)):
    return [{"id": a["id"], "title": a["title"], "company": a["company"], "match_score": a["match_score"],
             "created_at": iso(a["created_at"]), "analysis": a["result"].model_dump()}
            for a in database.list_analyses(user["id"])]


@app.delete("/api/analyses/{analysis_id}")
def delete_analysis(analysis_id: int, user: dict = Depends(current_user)):
    database.delete_analysis(user["id"], analysis_id)
    return {}


# --- Skills --------------------------------------------------------------------------------

def skills_out(user: dict) -> dict:
    analyses = database.list_analyses(user["id"])
    row = database.latest_coaching_report(user["id"])
    out = {"analyses_count": len(analyses), "min_analyses": coaching.MIN_ANALYSES,
           "new_since_report": len({a["id"] for a in analyses} - set(row["analysis_ids"])) if row else len(analyses),
           "report": None}
    if not row:
        return out
    report = coaching.SkillReport.model_validate_json(row["result_json"])
    by_id = {a["id"]: a for a in analyses}
    role_ids = database.role_ids_by_analysis(user["id"])

    def role_ref(idx):
        analysis_id = row["analysis_ids"][idx]
        a = by_id.get(analysis_id)
        return {"title": a["title"] if a else "Role", "company": a["company"] if a else "",
                "role_id": role_ids.get(analysis_id)}

    out["report"] = {
        "summary": report.summary, "roles_count": len(row["analysis_ids"]), "created_at": iso(row["created_at"]),
        "themes": [{**t.model_dump(exclude={"role_indexes"}),
                    "roles": [role_ref(i) for i in t.role_indexes if 0 <= i < len(row["analysis_ids"])]}
                   for t in report.themes],
    }
    return out


@app.get("/api/skills")
def get_skills(user: dict = Depends(current_user)):
    return skills_out(user)


@app.post("/api/skills")
def build_skills(user: dict = Depends(current_user)):
    analyses = database.list_analyses(user["id"])
    if len(analyses) < coaching.MIN_ANALYSES:
        raise HTTPException(400, f"Analyze at least {coaching.MIN_ANALYSES} roles first.")
    limits.require(user, "coaching")

    def work(log):
        log(f"Finding patterns across {len(analyses)} roles…")
        u = database.get_user(user["id"])
        report = coaching.build_report([a["result"] for a in analyses], u["resume_text"], u["resume_pdf"])
        limits.use(u, "coaching")
        database.save_coaching_report(u["id"], [a["id"] for a in analyses], report.model_dump_json())
        return skills_out(user)
    return start_job(user, work)


# --- Connections ---------------------------------------------------------------------------

@app.put("/api/connections")
async def put_connections(file: UploadFile = File(...), user: dict = Depends(current_user)):
    try:
        people = referrals.parse_linkedin_csv(await file.read())
    except ValueError as e:
        raise HTTPException(400, str(e))
    database.replace_connections(user["id"], people)
    return {"count": len(people)}


@app.delete("/api/connections")
def delete_connections(user: dict = Depends(current_user)):
    database.delete_connections(user["id"])
    return {"count": 0}


@app.get("/api/health")
def health():
    return {"ok": True}
